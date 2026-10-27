# -*- coding: utf-8 -*-
"""RAW / R1 / R2 / R2' metrics."""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from scipy import stats

_ROOT = Path(__file__).resolve().parents[1]
_R1R2 = _ROOT / "r1r2"
for _p in (str(_R1R2), str(_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from core.study_common import rgb, load_constraints, NAMES
from evaluation.eval_q import readout_Q, vlc_exact
from evaluation.eval_spatial import (
    road_metrics, area_element_metrics, green_space_metrics, composition_ratios,
    ID_BUILDING, ID_PAVEMENT, ID_WATER,
)
from evaluation.eval_checks import run_all_checks

OUT120 = None
PIXPIX = None
CONS_ROOT = None
NEW_TEST = None
SA_ROOT = None
SAMPLES_CSV = None

STAGES_120 = ("RAW", "R1", "R2", "R2_prime")
STAGE_LABEL = {
    "RAW": "Raw SD output",
    "R1": "R1 output",
    "R2": "R2 output",
    "R2_prime": "R2′ output",
}
AREA_PREFIX = {"building": "bld", "pavement": "pav", "water": "wat"}
AREA_CLASSES = (("building", ID_BUILDING), ("pavement", ID_PAVEMENT), ("water", ID_WATER))

COMP_KEYS = ("GSR", "PAR", "BCR", "WBR", "RAR", "PVR", "SUM")
T4_KEYS = ("NP_P", "MaxCS_P", "CLUMPY_P", "NGC", "MaxCS_G")
T5_KEYS = (
    "DIVISION_bld", "SHAPE_AM_bld",
    "DIVISION_pav", "SHAPE_AM_pav",
    "DIVISION_wat", "SHAPE_AM_wat",
)


def view_to_folder(view_id: str) -> str:
    """Map 000010_r090 → 000010_90 ; 000010_r000 → 000010_r000 ; 000010_r180 → 000010_180."""
    if view_id.endswith("_r000"):
        return view_id
    if view_id.endswith("_r090"):
        return view_id.replace("_r090", "_90")
    if view_id.endswith("_r180"):
        return view_id.replace("_r180", "_180")
    if view_id.endswith("_r270"):
        return view_id.replace("_r270", "_270")
    return view_id


def resolve_paths_120(view_id: str) -> Dict[str, Path]:
    folder = view_to_folder(view_id)
    park = view_id.split("_")[0]
    rot = int(view_id.split("_r")[-1]) if "_r" in view_id else 0
    # constraints: r000 uses park folder name without suffix
    if rot == 0:
        cdir = CONS_ROOT / park
        adir = NEW_TEST / f"{park}_r000"
    else:
        cdir = CONS_ROOT / folder
        adir = NEW_TEST / folder
    raw_dir = OUT120 / "raw_output" / folder
    return dict(
        raw=raw_dir / "raw.png",
        r1=OUT120 / "r1_output" / f"{view_id}.png",
        r2=OUT120 / "r2_output" / f"{view_id}.png",
        r2p=OUT120 / "r2'_output" / f"{view_id}.png",
        condition=adir / "A.png",
        site=cdir / "S.png",
        protect=cdir / "P.png",
        urban=cdir / "U.png",
    )


def resolve_paths_pixpix(view_id: str) -> Dict[str, Path]:
    folder = view_to_folder(view_id)
    park = view_id.split("_")[0]
    rot = int(view_id.split("_r")[-1]) if "_r" in view_id else 0
    if rot == 0:
        cdir = CONS_ROOT / park
        adir = NEW_TEST / f"{park}_r000"
    else:
        cdir = CONS_ROOT / folder
        adir = NEW_TEST / folder
    return dict(
        raw=PIXPIX / "_work_up512" / f"{view_id}_raw512.png",
        r1=PIXPIX / "r1" / f"{view_id}.png",
        r2=PIXPIX / "r2" / f"{view_id}.png",
        condition=adir / "A.png",
        site=cdir / "S.png",
        protect=cdir / "P.png",
        urban=cdir / "U.png",
    )


def load_sa_points_for_park(park_id: str) -> List[dict]:
    csv_path = SA_ROOT / park_id / f"{park_id}_labels.csv"
    if not csv_path.is_file():
        return []
    rows = []
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            human = int(float(r["human"]))
            if human not in range(1, 7):
                continue
            rows.append(dict(
                case_id=park_id,
                point_id=r.get("point_id"),
                x=int(float(r["x"])),
                y=int(float(r["y"])),
                human=human,
                status="valid",
                class_name=NAMES.get(human, ""),
            ))
    return rows


def compute_sa(Q: np.ndarray, domain: np.ndarray, protect: np.ndarray,
               points: List[dict]) -> Dict[str, Any]:
    h, w = Q.shape
    decidable = 0
    correct = 0
    excluded = 0
    for p in points:
        x, y = p["x"], p["y"]
        if not (0 <= x < w and 0 <= y < h):
            excluded += 1
            continue
        on_site = bool(domain[y, x] or protect[y, x])
        if not on_site:
            excluded += 1
            continue
        decidable += 1
        if int(Q[y, x]) == int(p["human"]):
            correct += 1
    return dict(
        SA=(correct / decidable) if decidable else None,
        n_decidable=decidable,
        n_correct=correct,
        n_excluded=excluded,
        n_points=len(points),
    )


def metrics_for_image(image: np.ndarray, cons: dict, stage: str,
                      sa_points: Optional[List[dict]], do_sa: bool) -> dict:
    domain = cons["editable"]
    protect = cons["protect"]
    Q = readout_Q(image)
    out: Dict[str, Any] = dict(stage=stage)
    out["PA"] = float(vlc_exact(image, domain))
    out.update(composition_ratios(Q, domain))
    out.update(road_metrics(Q, domain))
    green = green_space_metrics(Q, domain)
    out["NP_P"] = green["NP_P"]
    out["MaxCS_P"] = green["MaxCS_P"] if green["MaxCS_P"] is not None else (0 if green["NP_P"] == 0 else None)
    out["CLUMPY_P"] = green["CLUMPY_P"]
    out["NGC"] = green["NGC"]
    out["MaxCS_G"] = green["MaxCS_lawn"] if green["MaxCS_lawn"] is not None else (0 if green["NGC"] == 0 else None)
    out["clumpy_status"] = green["clumpy_status"]
    for cname, cid in AREA_CLASSES:
        am = area_element_metrics(Q, domain, cid, cname)
        pref = AREA_PREFIX[cname]
        out[f"DIVISION_{pref}"] = am["DIVISION_c"]
        out[f"DIVISION_{pref}_absent"] = am["DIVISION_absent"]
        out[f"SHAPE_AM_{pref}"] = am["SHAPE_AM_c"]
        out[f"NP_{pref}"] = am["NP_c"]
        out[f"A_{pref}"] = am["A_c"]
    out["domain_px"] = int(domain.sum())
    if do_sa and sa_points is not None:
        sa = compute_sa(Q, domain, protect, sa_points)
        out.update(sa)
    else:
        out.update(SA=None, n_decidable=0, n_correct=0, n_excluded=0, n_points=0)
    return out


def _mean_std(vals: List[float]) -> Tuple[Optional[str], int]:
    a = [v for v in vals if v is not None and v == v]
    if not a:
        return None, 0
    arr = np.asarray(a, dtype=float)
    return f"{arr.mean():.3f}±{arr.std(ddof=0):.3f}", len(a)


def _median_iqr(vals: List[float], *, as_int: bool = False) -> Tuple[Optional[str], int]:
    a = [v for v in vals if v is not None and v == v]
    if not a:
        return None, 0
    arr = np.asarray(a, dtype=float)
    q1, med, q3 = np.percentile(arr, [25, 50, 75])
    if as_int:
        return f"{int(round(med))} ({int(round(q1))}, {int(round(q3))})", len(a)
    return f"{med:.3f} ({q1:.3f}, {q3:.3f})", len(a)


def _sci_mean_std(vals: List[float]) -> Tuple[Optional[str], int]:
    a = [v for v in vals if v is not None and v == v]
    if not a:
        return None, 0
    arr = np.asarray(a, dtype=float)
    m, s = float(arr.mean()), float(arr.std(ddof=0))
    if abs(m) < 1e-3 and abs(s) < 1e-3:
        return f"{m:.2E}±{s:.2E}", len(a)
    return f"{m:.3f}±{s:.3f}", len(a)


def park_means(rows: List[dict], key: str) -> Dict[str, Dict[str, float]]:
    """park -> stage -> mean over views (NaN-safe)."""
    bucket: Dict[str, Dict[str, List[float]]] = defaultdict(lambda: defaultdict(list))
    for r in rows:
        v = r.get(key)
        if v is None or v != v:
            continue
        bucket[r["park_id"]][r["stage"]].append(float(v))
    out: Dict[str, Dict[str, float]] = {}
    for park, stmap in bucket.items():
        out[park] = {st: float(np.mean(vs)) for st, vs in stmap.items()}
    return out


def paired_tests(rows: List[dict], key: str, pairs: List[Tuple[str, str]]) -> dict:
    pm = park_means(rows, key)
    parks = sorted(pm)
    result = {}
    # Friedman on parks that have all four stages
    stages = list(STAGES_120)
    mats = []
    parks_f = []
    for p in parks:
        if all(st in pm[p] for st in stages):
            mats.append([pm[p][st] for st in stages])
            parks_f.append(p)
    if len(mats) >= 3:
        cols = list(zip(*mats))
        try:
            stat, pval = stats.friedmanchisquare(*cols)
            result["friedman"] = dict(stat=float(stat), p=float(pval), n_parks=len(mats))
        except Exception as e:
            result["friedman"] = dict(error=str(e), n_parks=len(mats))
    else:
        result["friedman"] = dict(n_parks=len(mats), note="insufficient")
    result["pairs"] = {}
    for a, b in pairs:
        diffs = []
        for p in parks:
            if a in pm[p] and b in pm[p]:
                diffs.append(pm[p][a] - pm[p][b])
        if len(diffs) < 3:
            result["pairs"][f"{a}_vs_{b}"] = dict(n=len(diffs))
            continue
        # Wilcoxon requires non-zero diffs; zero-diff handling
        arr = np.asarray(diffs, dtype=float)
        if np.allclose(arr, 0):
            result["pairs"][f"{a}_vs_{b}"] = dict(n=len(diffs), p=1.0, note="all_zero")
            continue
        try:
            w = stats.wilcoxon(arr, zero_method="wilcox", alternative="two-sided")
            result["pairs"][f"{a}_vs_{b}"] = dict(
                n=len(diffs), stat=float(w.statistic), p=float(w.pvalue),
                median_diff=float(np.median(arr)),
            )
        except Exception as e:
            result["pairs"][f"{a}_vs_{b}"] = dict(n=len(diffs), error=str(e))
    return result


def summarize_tables(rows: List[dict]) -> Dict[str, Any]:
    by_stage = {st: [r for r in rows if r["stage"] == st] for st in STAGES_120}
    # also allow subset stages
    present = [st for st in STAGES_120 if by_stage[st]]
    tables = {}

    # Table 1
    t1 = {}
    for st in present:
        pa_s, n_pa = _mean_std([r["PA"] for r in by_stage[st]])
        sa_rows = [r for r in by_stage[st] if int(r.get("rotation", -1)) == 0]
        sa_s, n_sa = _mean_std([r["SA"] for r in sa_rows])
        t1[st] = dict(PA=pa_s, SA=sa_s, n_PA=n_pa, n_SA=n_sa)
    tables["Table1"] = t1

    # Table 2
    t2 = {}
    for st in present:
        cell = {}
        for k in COMP_KEYS:
            if k == "RAR":
                s, n = _sci_mean_std([r[k] for r in by_stage[st]])
            else:
                s, n = _mean_std([r[k] for r in by_stage[st]])
            cell[k] = s
            cell[f"n_{k}"] = n
        t2[st] = cell
    tables["Table2"] = t2

    # Table 3
    t3 = {}
    for st in present:
        lcr_s, n_l = _median_iqr([r["LCR"] for r in by_stage[st]])
        rer_s, n_r = _median_iqr([r["RER"] for r in by_stage[st]])
        t3[st] = dict(LCR=lcr_s, RER=rer_s, n_LCR=n_l, n_RER=n_r)
    tables["Table3"] = t3

    # Table 4
    t4 = {}
    for st in present:
        cell = {}
        for k in T4_KEYS:
            as_int = k in ("NP_P", "NGC", "MaxCS_P", "MaxCS_G")
            s, n = _median_iqr([r[k] for r in by_stage[st]], as_int=as_int)
            cell[k] = s
            cell[f"n_{k}"] = n
        t4[st] = cell
    tables["Table4"] = t4

    # Table 5
    t5 = {}
    for st in present:
        cell = {}
        for k in T5_KEYS:
            s, n = _median_iqr([r[k] for r in by_stage[st]])
            cell[k] = s
            cell[f"n_{k}"] = n
        t5[st] = cell
    tables["Table5"] = t5

    pairs = [("R1", "RAW"), ("R2", "R1"), ("R2", "R2_prime"), ("R2_prime", "RAW")]
    tests = {}
    for key in ("PA", "SA", "SUM", "LCR", "RER", "NGC", "MaxCS_G", "NP_P", "MaxCS_P",
                "CLUMPY_P", *T5_KEYS):
        # SA only r000
        src = rows if key != "SA" else [r for r in rows if int(r.get("rotation", -1)) == 0]
        tests[key] = paired_tests(src, key, pairs)
    tables["tests"] = tests
    return tables


def write_excel(path: Path, per_view: List[dict], tables: dict, meta: dict):
    wb = Workbook()
    # per-view
    ws = wb.active
    ws.title = "per_view"
    keys = [
        "dataset", "view_id", "park_id", "rotation", "stage",
        "PA", "SA", "n_decidable", "n_correct",
        *COMP_KEYS, "LCR", "RER", "road_px",
        *T4_KEYS, *T5_KEYS, "domain_px",
    ]
    ws.append(keys)
    for r in per_view:
        ws.append([r.get(k) for k in keys])

    def add_summary_sheet(name: str, headers: List[str], col_keys: List[str], table: dict):
        w = wb.create_sheet(name)
        w.append(["Stage", *headers])
        for st in STAGES_120:
            if st not in table:
                continue
            row = [STAGE_LABEL[st]]
            for k in col_keys:
                row.append(table[st].get(k))
            w.append(row)
        for col in w.columns:
            for cell in col:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    add_summary_sheet("Table1_PA_SA", ["PA", "SA"], ["PA", "SA"], tables["Table1"])
    add_summary_sheet("Table2_Composition", list(COMP_KEYS), list(COMP_KEYS), tables["Table2"])
    add_summary_sheet("Table3_Road", ["LCR", "RER"], ["LCR", "RER"], tables["Table3"])
    add_summary_sheet("Table4_Green", list(T4_KEYS), list(T4_KEYS), tables["Table4"])
    add_summary_sheet(
        "Table5_AreaElements",
        ["DIVISION_bld", "SHAPE_AM_bld", "DIVISION_pav", "SHAPE_AM_pav",
         "DIVISION_wat", "SHAPE_AM_wat"],
        list(T5_KEYS), tables["Table5"],
    )

    w = wb.create_sheet("Notes_tests")
    w.append(["metric", "friedman_p", "n_parks", "pair", "wilcoxon_p", "n", "median_diff"])
    for metric, tr in tables.get("tests", {}).items():
        fr = tr.get("friedman", {})
        pairs = tr.get("pairs", {})
        if not pairs:
            w.append([metric, fr.get("p"), fr.get("n_parks"), "", "", "", ""])
            continue
        first = True
        for pname, pr in pairs.items():
            w.append([
                metric if first else "",
                fr.get("p") if first else "",
                fr.get("n_parks") if first else "",
                pname, pr.get("p"), pr.get("n"), pr.get("median_diff"),
            ])
            first = False

    w = wb.create_sheet("meta")
    w.append(["key", "value"])
    for k, v in meta.items():
        w.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def run_dataset(name: str, view_ids: List[str], path_fn, stages: Tuple[str, ...],
                sa_cache: Dict[str, List[dict]]) -> List[dict]:
    rows = []
    for i, view_id in enumerate(view_ids):
        park = view_id.split("_")[0]
        rot = int(view_id.split("_r")[-1])
        paths = path_fn(view_id)
        for k in ("condition", "site", "protect", "urban"):
            if not paths[k].is_file():
                raise FileNotFoundError(f"{view_id}: missing {k}={paths[k]}")
        sample = dict(
            view_id=view_id, park_id=park, rotation=str(rot),
            condition_path=str(paths["condition"]),
            site_path=str(paths["site"]),
            protect_path=str(paths["protect"]),
            urban_path=str(paths["urban"]),
        )
        cons = load_constraints(sample)
        sa_points = sa_cache.get(park, [])
        do_sa = (rot == 0) and bool(sa_points)
        stage_files = {
            "RAW": paths.get("raw"),
            "R1": paths.get("r1"),
            "R2": paths.get("r2"),
            "R2_prime": paths.get("r2p"),
        }
        for st in stages:
            fp = stage_files.get(st)
            if fp is None or not fp.is_file():
                print(f"  SKIP {view_id} {st}: missing {fp}")
                continue
            img = rgb(fp)
            if img.shape[:2] != cons["a"].shape[:2]:
                raise ValueError(f"{view_id} {st} size {img.shape[:2]} vs A {cons['a'].shape[:2]}")
            m = metrics_for_image(img, cons, st, sa_points, do_sa)
            m.update(dataset=name, view_id=view_id, park_id=park, rotation=rot,
                     image_path=str(fp))
            rows.append(m)
        if (i + 1) % 20 == 0 or (i + 1) == len(view_ids):
            print(f"  [{name}] {i+1}/{len(view_ids)}")
    return rows


def main():
    global OUT120, PIXPIX, CONS_ROOT, NEW_TEST, SA_ROOT, SAMPLES_CSV
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-root", type=Path, required=True)
    ap.add_argument("--samples", type=Path, default=None)
    ap.add_argument("--out", type=Path, default=Path("metrics_out"))
    ap.add_argument("--skip-checks", action="store_true")
    ap.add_argument("--only-pixpix", action="store_true")
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()
    root = Path(args.data_root)
    OUT120 = root / "outputs"
    PIXPIX = root / "pixpix"
    CONS_ROOT = root / "constraints"
    NEW_TEST = root / "conditions"
    SA_ROOT = root / "sa"
    SAMPLES_CSV = Path(args.samples) if args.samples else root / "samples.csv"
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    if not args.skip_checks:
        chk = run_all_checks()
        (out_dir / "eval_checks.json").write_text(
            json.dumps(chk, ensure_ascii=False, indent=2), encoding="utf-8")
        print("CHECKS", chk)

    # SA cache
    sa_cache = {}
    for d in sorted(p for p in SA_ROOT.iterdir() if p.is_dir()):
        sa_cache[d.name] = load_sa_points_for_park(d.name)
    print(f"SA parks={len(sa_cache)} points={sum(len(v) for v in sa_cache.values())}")

    with SAMPLES_CSV.open(encoding="utf-8-sig", newline="") as f:
        samples = list(csv.DictReader(f))
    view_ids_120 = [r["view_id"] for r in samples]
    if args.limit:
        view_ids_120 = view_ids_120[: args.limit]

    all_rows = []
    if not args.only_pixpix:
        print("RUN 120")
        rows120 = run_dataset("120_test", view_ids_120, resolve_paths_120, STAGES_120, sa_cache)
        all_rows.extend(rows120)
        tables120 = summarize_tables(rows120)
        write_excel(
            out_dir / "metrics_120.xlsx", rows120, tables120,
            meta=dict(n_views=len(view_ids_120), n_rows=len(rows120),
                      sa_parks=len(sa_cache), stages=list(STAGES_120)),
        )
        (out_dir / "tables_120.json").write_text(
            json.dumps(tables120, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print("Wrote metrics_120.xlsx")

    # pixpix 4
    pix_views = sorted(p.stem for p in (PIXPIX / "r1").glob("*.png"))
    print("RUN pixpix", pix_views)
    rows_px = run_dataset(
        "pixpix_4", pix_views, resolve_paths_pixpix, ("RAW", "R1", "R2"), sa_cache)
    all_rows.extend(rows_px)
    # summarize with available stages only
    tables_px = summarize_tables(rows_px)
    write_excel(
        out_dir / "metrics_pixpix4.xlsx", rows_px, tables_px,
        meta=dict(n_views=len(pix_views), n_rows=len(rows_px), stages=["RAW", "R1", "R2"]),
    )
    (out_dir / "tables_pixpix4.json").write_text(
        json.dumps(tables_px, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    # combined workbook convenience
    if all_rows:
        tables_all = summarize_tables([r for r in all_rows if r["dataset"] == "120_test"] or all_rows)
        write_excel(out_dir / "metrics_all_detail.xlsx", all_rows, tables_all,
                    meta=dict(note="per_view includes 120 + pixpix4"))

    print("DONE", out_dir)


if __name__ == "__main__":
    main()
